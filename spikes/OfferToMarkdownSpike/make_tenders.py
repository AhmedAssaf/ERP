"""Generate three fictional tenders, three offers each, and an answer key per tender.

Output, samples/tenders/<tender>/:
  rfp.docx and/or rfp.pdf
  offers/<vendor>/technical.<docx|pdf>   as the vendor would submit it (docx, Word PDF, or image-only scan)
  offers/<vendor>/financial.<same>       the separate financial envelope
  answer-key.json                        expected verdict per requirement and the planted flags
Needs Microsoft Word for the PDF export.
"""
import json
import shutil
from pathlib import Path

import win32com.client
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

from make_samples import rasterize, set_rtl, word_to_pdf
from tenders_data import AR_FOOTER, EN_FOOTER, TENDERS

OUT = Path(__file__).parent / "samples" / "tenders"


def build_docx(blocks, path, rtl):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:cs"), "Arial")
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = AR_FOOTER if rtl else EN_FOOTER
    if rtl:
        set_rtl(footer)
    for block in blocks:
        kind = block[0]
        if kind == "br":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue
        if kind == "t":
            rows = block[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            if rtl:
                t._tbl.tblPr.append(t._tbl.tblPr.makeelement(qn("w:bidiVisual"), {}))
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    cell = t.cell(r, c).paragraphs[0]
                    cell.text = text
                    if r == 0:
                        cell.runs[0].bold = True
                    if rtl:
                        set_rtl(cell)
            doc.add_paragraph()
            continue
        p = doc.add_heading(block[1], level=1) if kind == "h" else doc.add_paragraph(block[1])
        if rtl:
            set_rtl(p)
    doc.save(path)


def emit(word, blocks, stem, fmt, rtl):
    """Write stem.docx, then keep it or replace it with a PDF or an image-only scan."""
    docx = stem.with_suffix(".docx")
    build_docx(blocks, docx, rtl)
    if fmt == "docx":
        return docx
    pdf = stem.with_suffix(".pdf")
    word_to_pdf(word, docx, pdf)
    docx.unlink()
    if fmt == "scanned":
        text_pdf = stem.with_name(stem.name + "-text.pdf")
        pdf.rename(text_pdf)
        rasterize(text_pdf, pdf, dpi=150)
        text_pdf.unlink()
    return pdf


def requirements(rfp):
    """Requirement ID to text, from the RFP's requirement tables."""
    reqs = {}
    for block in rfp:
        if block[0] == "t":
            for row in block[1][1:]:
                if len(row) == 2 and row[0][:2] in ("M-", "T-"):
                    reqs[row[0]] = row[1]
    return reqs


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        for t in TENDERS:
            folder = OUT / t["id"]
            folder.mkdir(parents=True)
            for fmt in sorted(t["rfp_formats"], key=lambda f: f == "docx"):  # PDF export removes its docx, so docx goes last
                emit(word, t["rfp"], folder / "rfp", fmt, t["rtl"])
            key = {"tender": t["id"], "requirements": requirements(t["rfp"]), "offers": {}}
            for o in t["offers"]:
                vdir = folder / "offers" / o["id"]
                vdir.mkdir(parents=True)
                tech = emit(word, o["tech"], vdir / "technical", o["format"], t["rtl"])
                fin = emit(word, o["fin"], vdir / "financial", o["format"], t["rtl"])
                key["offers"][o["id"]] = {"name": o["name"], "format": o["format"],
                                          "technical": tech.relative_to(folder).as_posix(),
                                          "financial": fin.relative_to(folder).as_posix(),
                                          "expected": o["expected"], "flags": o["flags"]}
                print(f"{t['id']:20} {o['id']:16} {o['format']}", flush=True)
            (folder / "answer-key.json").write_text(json.dumps(key, ensure_ascii=False, indent=2), encoding="utf-8")
    finally:
        word.Quit()


if __name__ == "__main__":
    main()
