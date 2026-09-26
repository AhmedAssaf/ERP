"""Build sample offers with known text for spike W-22.

Writes to samples/generated/:
  offer-ar.docx          Arabic technical offer with English runs, a table, 3 pages
  offer-ar.pdf           the same, exported by Microsoft Word (text PDF, like most real offers)
  offer-ar-scanned.pdf   the same pages rendered to 200 dpi images (no text layer)
  offer-en.docx / .pdf   English-only control
  truth.json             the source text per file and page, in logical order
"""
import json
from pathlib import Path

import pymupdf as fitz  # installed with pymupdf4llm
import win32com.client
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).parent / "samples" / "generated"

AR_PAGES = [
    {
        "heading": "العرض الفني لمنافسة توريد وتركيب شبكة البيانات",
        "paragraphs": [
            "تتشرف شركة الأفق للحلول التقنية بتقديم عرضها الفني للمنافسة رقم 2026-114 الخاصة بتوريد وتركيب شبكة البيانات للمقر الرئيسي في الرياض.",
            "نلتزم بتوريد محولات Cisco Catalyst 9300 عدد 24 وحدة مع رخص DNA Advantage لمدة ثلاث سنوات.",
            "مدة التنفيذ 90 يوما من تاريخ استلام أمر الشراء، وتشمل التركيب والاختبار والتدريب.",
        ],
    },
    {
        "heading": "جدول المطابقة للمتطلبات الفنية",
        "table": [
            ["رقم المتطلب", "المتطلب", "الحالة", "المرجع في العرض"],
            ["TR-01", "دعم معيار IEEE 802.3bt لتغذية الأجهزة", "مطابق", "صفحة 2"],
            ["TR-02", "ضمان لمدة خمس سنوات على الأجهزة", "مطابق جزئيا", "صفحة 3"],
            ["TR-03", "دعم فني على مدار الساعة 24/7 باللغة العربية", "مطابق", "صفحة 3"],
            ["TR-04", "شهادة ISO 27001 سارية للمورد", "غير مطابق", "لا يوجد"],
        ],
        "paragraphs": [
            "ملاحظة: الضمان المقدم ثلاث سنوات من الشركة المصنعة وسنتان إضافيتان من الشركة.",
        ],
    },
    {
        "heading": "فريق العمل والتدريب",
        "paragraphs": [
            "يقود المشروع مهندس معتمد بشهادة CCNP Enterprise وخبرة تزيد على عشر سنوات.",
            "يشمل التدريب 5 أيام لعدد 8 موظفين من إدارة تقنية المعلومات في مقر العميل.",
            "نتعهد بتسليم المخططات النهائية بصيغة PDF و Visio خلال 14 يوما من الاستلام الابتدائي.",
        ],
    },
]

EN_PAGES = [
    {
        "heading": "Technical Offer: Data Network Supply and Installation",
        "paragraphs": [
            "Horizon Technical Solutions submits its technical offer for tender 2026-114 for the head office data network in Riyadh.",
            "We will supply 24 Cisco Catalyst 9300 switches with DNA Advantage licences for three years.",
        ],
    },
    {
        "heading": "Compliance Matrix",
        "table": [
            ["Req", "Requirement", "Status", "Offer reference"],
            ["TR-01", "IEEE 802.3bt power over Ethernet", "Compliant", "Page 2"],
            ["TR-04", "Valid ISO 27001 certificate", "Not compliant", "None"],
        ],
        "paragraphs": ["Note: warranty is three years from the manufacturer plus two from us."],
    },
]


def set_rtl(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.makeelement(qn("w:bidi"), {})
    ppr.append(bidi)


def build_docx(pages, path, rtl):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:cs"), "Arial")
    for i, page in enumerate(pages):
        h = doc.add_heading(page["heading"], level=1)
        if rtl:
            set_rtl(h)
        if "table" in page:
            t = doc.add_table(rows=len(page["table"]), cols=len(page["table"][0]))
            t.style = "Table Grid"
            if rtl:
                tblpr = t._tbl.tblPr
                tblpr.append(tblpr.makeelement(qn("w:bidiVisual"), {}))
            for r, row in enumerate(page["table"]):
                for c, text in enumerate(row):
                    cell_p = t.cell(r, c).paragraphs[0]
                    cell_p.text = text
                    if rtl:
                        set_rtl(cell_p)
        for text in page["paragraphs"]:
            p = doc.add_paragraph(text)
            if rtl:
                set_rtl(p)
        if i < len(pages) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(path)


def page_text(page):
    parts = [page["heading"]]
    for row in page.get("table", []):
        parts.extend(row)
    parts.extend(page["paragraphs"])
    return "\n".join(parts)


def word_to_pdf(word, docx_path, pdf_path):
    d = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
    d.SaveAs2(str(pdf_path.resolve()), FileFormat=17)  # wdFormatPDF
    d.Close(False)


def rasterize(pdf_path, out_path, dpi=200):
    src = fitz.open(pdf_path)
    dst = fitz.open()
    for page in src:
        pix = page.get_pixmap(dpi=dpi)
        new = dst.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=pix.tobytes("jpeg", jpg_quality=80))
    dst.save(out_path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(AR_PAGES, OUT / "offer-ar.docx", rtl=True)
    build_docx(EN_PAGES, OUT / "offer-en.docx", rtl=False)

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        word_to_pdf(word, OUT / "offer-ar.docx", OUT / "offer-ar.pdf")
        word_to_pdf(word, OUT / "offer-en.docx", OUT / "offer-en.pdf")
    finally:
        word.Quit()

    rasterize(OUT / "offer-ar.pdf", OUT / "offer-ar-scanned.pdf")

    ar = [page_text(p) for p in AR_PAGES]
    en = [page_text(p) for p in EN_PAGES]
    truth = {
        "offer-ar.docx": ar,
        "offer-ar.pdf": ar,
        "offer-ar-scanned.pdf": ar,
        "offer-en.docx": en,
        "offer-en.pdf": en,
    }
    (OUT / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
    print("wrote", ", ".join(sorted(truth)))


if __name__ == "__main__":
    main()
